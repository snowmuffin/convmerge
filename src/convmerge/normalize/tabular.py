"""CSV, TSV and Excel files → JSONL, one object per row keyed by the header.

Spreadsheet exports are a common way instruction data arrives (``instruction,
input,output`` columns). The first row is the header; every later row becomes
an object of strings. Cells that span lines (quoted) are kept whole, a UTF-8
byte order mark is ignored, empty rows are skipped, and header cells that are
empty or repeated get unique names (``column_3``, ``output_2``) so no value is
lost. Values stay strings; ``convert`` reads JSON-looking strings such as a
serialized ``messages`` column itself.

Excel workbooks (``.xlsx`` / ``.xlsm``, the ``xlsx`` extra) follow the same
rules on one sheet (the first, or ``sheet``), with cell values written as
text: numbers with Excel's 15 significant digits (``3``, ``0.3``), dates
as ISO 8601, booleans as ``TRUE`` / ``FALSE``. A formula cell holds the value
Excel saved with the file; files written by a program and never opened in
Excel have none, and those cells come out empty (counted in a warning).
"""

from __future__ import annotations

import csv
import datetime
import json
import logging
import sys
from collections.abc import Iterable, Iterator
from contextlib import ExitStack
from pathlib import Path
from typing import Any, TextIO

from convmerge._output import atomic_text_writer
from convmerge.io import refuse_overwrite

TABLE_EXTENSIONS: tuple[str, ...] = (".csv", ".tsv")
XLSX_EXTENSIONS: tuple[str, ...] = (".xlsx", ".xlsm")

logger = logging.getLogger(__name__)


def table_to_jsonl(src: str | Path, dst: str | Path, *, encoding: str = "utf-8") -> int:
    """Rewrite a ``.csv`` (comma) or ``.tsv`` (tab) file as JSONL; return rows written."""
    refuse_overwrite([src], [dst])
    src_p, dst_p = Path(src), Path(dst)
    delimiter = "\t" if src_p.suffix.lower() == ".tsv" else ","
    _allow_long_fields()
    dst_p.parent.mkdir(parents=True, exist_ok=True)
    enc = "utf-8-sig" if encoding.lower().replace("_", "-") in ("utf-8", "utf8") else encoding
    with atomic_text_writer(dst_p) as fout, src_p.open(encoding=enc, newline="") as fin:
        reader = csv.reader(fin, delimiter=delimiter, strict=True)
        try:
            written = _write_rows(reader, fout)
        except csv.Error as e:
            raise ValueError(f"{src_p}: line {reader.line_num}: {e}") from None
    return written


def xlsx_to_jsonl(src: str | Path, dst: str | Path, *, sheet: str | None = None) -> int:
    """Rewrite one sheet of an ``.xlsx`` workbook as JSONL; return rows written.

    ``sheet`` names the sheet (default: the first). Needs ``openpyxl``
    (``pip install "convmerge[xlsx]"``).
    """
    refuse_overwrite([src], [dst])
    src_p, dst_p = Path(src), Path(dst)
    if src_p.suffix.lower() == ".xls":
        raise ValueError(f"{src_p}: .xls (Excel 97-2003) is not read; save it as .xlsx or .csv")
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            'reading .xlsx files needs openpyxl: pip install "convmerge[xlsx]" (or [all])'
        ) from None
    with atomic_text_writer(dst_p) as fout, ExitStack() as readers:
        try:
            values = openpyxl.load_workbook(src_p, read_only=True, data_only=True)
            readers.callback(values.close)
            formulas = openpyxl.load_workbook(src_p, read_only=True, data_only=False)
            readers.callback(formulas.close)
        except Exception as e:  # noqa: BLE001 - zip, XML, and openpyxl's own errors
            raise ValueError(
                f"{src_p}: not a readable .xlsx workbook ({type(e).__name__}: {e})"
            ) from None
        name = sheet if sheet is not None else values.sheetnames[0]
        if name not in values.sheetnames:
            raise ValueError(f"{src_p}: no sheet {name!r}; sheets: {', '.join(values.sheetnames)}")
        missing = [0]
        rows = _xlsx_rows(values[name], formulas[name], missing)
        written = _write_rows(rows, fout)
        if missing[0]:
            logger.warning(
                "%s: %d formula cell(s) have no saved value and were left empty; "
                "open and save the file in Excel (or LibreOffice) to store the values",
                src_p, missing[0],
            )  # fmt: skip
    return written


def _xlsx_rows(values: Any, formulas: Any, missing: list[int]) -> Iterator[list[str]]:
    for vrow, frow in zip(values.iter_rows(values_only=True), formulas.iter_rows(values_only=True)):
        cells: list[str] = []
        for value, formula in zip(vrow, frow):
            if value is None and isinstance(formula, str) and formula.startswith("="):
                missing[0] += 1
            cells.append(_cell_text(value))
        while cells and not cells[-1]:
            cells.pop()  # sheets often span formatted but empty columns
        yield cells


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return format(value, ".15g")  # Excel keeps 15 significant digits
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def _write_rows(reader: Iterable[list[str]], fout: TextIO) -> int:
    written = 0
    rows = iter(reader)
    header = next(rows, None)
    if header is None:
        return 0
    names = _column_names(header)
    for row in rows:
        if not any(cell.strip() for cell in row):
            continue
        record = {names[i] if i < len(names) else f"column_{i + 1}": cell
                  for i, cell in enumerate(row)}  # fmt: skip
        for name in names[len(row) :]:
            record[name] = ""
        fout.write(json.dumps(record, ensure_ascii=False) + "\n")
        written += 1
    return written


def _column_names(header: list[str]) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for i, cell in enumerate(header):
        base = cell.strip() or f"column_{i + 1}"
        name, n = base, 1
        while name in seen:
            n += 1
            name = f"{base}_{n}"
        seen.add(name)
        names.append(name)
    return names


def _allow_long_fields() -> None:
    # The csv module rejects fields over 128 KiB by default; long answers exceed it.
    limit = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limit)
            return
        except OverflowError:
            limit //= 10
